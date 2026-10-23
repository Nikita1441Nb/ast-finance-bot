import io
import openpyxl
import xlrd
import requests
from pypdf import PdfReader
from database import get_db

def decode_text_bytes(file_bytes: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return file_bytes.decode(enc)[:20000]
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")[:20000]

def parse_excel_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows.append(f"=== Лист: {sheet} ===")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append(" | ".join(str(c) if c is not None else "" for c in row).strip())
    return "\n".join(rows)[:20000]

def parse_excel_xls(file_bytes: bytes) -> str:
    wb = xlrd.open_workbook(file_contents=file_bytes)
    rows = []
    for sheet_name in wb.sheet_names():
        sheet = wb.sheet_by_name(sheet_name)
        rows.append(f"=== Лист: {sheet_name} ===")
        for row_idx in range(sheet.nrows):
            row_vals = sheet.row_values(row_idx)
            if any(v != "" and v is not None for v in row_vals):
                rows.append(" | ".join(str(v) if v != "" else "" for v in row_vals).strip())
    return "\n".join(rows)[:20000]

def parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text = []
    for idx, page in enumerate(reader.pages):
        t = page.extract_text()
        if t:
            text.append(f"--- Стр {idx+1} ---\n{t}")
    return "\n".join(text)[:20000]

def generate_excel_bytes() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Операции"
    ws1.append(["ID", "Дата записи", "Сотрудник", "Дата док.", "Тип", "Юрлицо", "Контрагент", "Вид", "Сумма (₸)", "Детали"])

    ws2 = wb.create_sheet(title="Сводки")
    ws2.append(["ID", "Дата записи", "Сотрудник", "Заголовок", "Описание карточки"])

    with get_db() as conn:
        rows = conn.execute(
            """SELECT id, created_at, added_by_name, doc_date, doc_type, company, counterparty, operation_type, amount, details
               FROM transactions ORDER BY id DESC"""
        ).fetchall()
        for r in rows:
            ws1.append(list(r))

        sum_rows = conn.execute(
            "SELECT id, created_at, added_by_name, title, formatted_summary FROM summaries ORDER BY id DESC"
        ).fetchall()
        for sr in sum_rows:
            ws2.append(list(sr))

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio

def get_weather_forecast(city: str) -> str:
    try:
        resp = requests.get(f"https://wttr.in/{city}?format=3", timeout=10)
        if resp.status_code == 200:
            return resp.text.strip()
        return f"Погода для {city} временно недоступна."
    except Exception as e:
        return f"Ошибка получения погоды: {e}"