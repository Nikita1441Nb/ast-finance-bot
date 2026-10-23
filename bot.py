import os
import sys
import io
import re
import json
import sqlite3
import base64
import asyncio
import logging
from datetime import datetime
from zoneinfo import ZoneInfo
from contextlib import contextmanager

TIMEZONE_ALMATY = ZoneInfo("Asia/Almaty")

def get_current_almaty_time() -> datetime:
    """Возвращает актуальное системное время в Asia/Almaty (UTC+5)."""
    return datetime.now(TIMEZONE_ALMATY)
from logging.handlers import RotatingFileHandler
from dotenv import load_dotenv
import openpyxl
import xlrd
from pypdf import PdfReader

from telegram import (
    Update,
    ReplyKeyboardMarkup,
    KeyboardButton,
    InlineKeyboardMarkup,
    InlineKeyboardButton
)
from telegram.error import BadRequest
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    filters,
    ContextTypes
)
from anthropic import AsyncAnthropic

BASE_DIR = r"C:\TelegramBot"
os.makedirs(BASE_DIR, exist_ok=True)
load_dotenv(os.path.join(BASE_DIR, ".env"))

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")
CLAUDE_MODEL = os.getenv("CLAUDE_MODEL", "claude-haiku-4-5-20251001")

USERS_FILE = os.path.join(BASE_DIR, "users.txt")
DB_PATH = os.path.join(BASE_DIR, "bot.db")
DATA_DIR = os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)

log_file = os.path.join(BASE_DIR, "bot.log")
logging.basicConfig(
    format="%(asctime)s [%(levelname)s] (%(name)s): %(message)s",
    level=logging.INFO,
    handlers=[
        RotatingFileHandler(log_file, maxBytes=5 * 1024 * 1024, backupCount=5, encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)
logging.getLogger("httpx").setLevel(logging.WARNING)
logger = logging.getLogger("ASTGroupBot")

claude_client = AsyncAnthropic(api_key=ANTHROPIC_API_KEY, max_retries=3)
SUPER_ADMINS = ["lexaengineer1441", "ast_aidar"]


@contextmanager
def get_db():
    conn = sqlite3.connect(DB_PATH, timeout=30.0)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA journal_mode=WAL;")
        conn.execute("PRAGMA synchronous=NORMAL;")
        conn.execute("PRAGMA busy_timeout=30000;")
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def init_db():
    with get_db() as conn:
        c = conn.cursor()
        c.execute("""
            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                added_by_username TEXT,
                added_by_name TEXT,
                doc_date TEXT,
                doc_type TEXT,
                company TEXT,
                counterparty TEXT,
                operation_type TEXT,
                amount REAL DEFAULT 0,
                details TEXT,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS summaries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                added_by_name TEXT,
                title TEXT,
                formatted_summary TEXT,
                raw_json TEXT,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS chat_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                role TEXT,
                content TEXT,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS drafts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                uinfo_json TEXT,
                data_json TEXT,
                waiting_edit INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS user_settings (
                user_id INTEGER PRIMARY KEY,
                mode TEXT
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS user_rules (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                rule_text TEXT,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS user_states (
                user_id INTEGER PRIMARY KEY,
                step TEXT,
                temp_name TEXT,
                temp_username TEXT,
                temp_target TEXT,
                temp_remind_text TEXT
            )
        """)
        # Автоматическая миграция новых колонок для существующей таблицы user_states
        cols = [r["name"] for r in conn.execute("PRAGMA table_info(user_states)").fetchall()]
        if "temp_target" not in cols:
            conn.execute("ALTER TABLE user_states ADD COLUMN temp_target TEXT;")
        if "temp_remind_text" not in cols:
            conn.execute("ALTER TABLE user_states ADD COLUMN temp_remind_text TEXT;")
        c.execute("""
            CREATE TABLE IF NOT EXISTS telegram_users (
                user_id INTEGER PRIMARY KEY,
                username TEXT,
                first_name TEXT
            )
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS reminders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                creator_id INTEGER,
                creator_name TEXT,
                target_user TEXT,
                remind_text TEXT,
                remind_time TEXT,
                is_sent INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT (datetime('now', 'localtime'))
            )
        """)


def get_user_info(username: str):
    uname_clean = (username or "").lstrip("@").strip().lower()
    if not os.path.exists(USERS_FILE):
        with open(USERS_FILE, "w", encoding="utf-8") as f:
            f.write("Никита : @Lexaengineer1441 : Админ, Разработчик\n")
            f.write("Айдар : @AST_Aidar : Админ, Руководитель\n")

    if uname_clean in SUPER_ADMINS:
        return {"name": "Администратор", "username": uname_clean, "desc": "Главный админ", "is_admin": True}

    with open(USERS_FILE, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or ":" not in line:
                continue
            parts = [p.strip() for p in line.split(":", 2)]
            if len(parts) >= 2:
                name = parts[0]
                file_uname = parts[1].lstrip("@").strip().lower()
                desc = parts[2] if len(parts) > 2 else "Сотрудник"
                if file_uname == uname_clean:
                    is_admin = ("админ" in desc.lower()) or (uname_clean in SUPER_ADMINS)
                    return {"name": name, "username": uname_clean, "desc": desc, "is_admin": is_admin}
    return None



def get_user_state(user_id: int):
    with get_db() as conn:
        row = conn.execute("SELECT step, temp_name, temp_username, temp_target, temp_remind_text FROM user_states WHERE user_id = ?", (user_id,)).fetchone()
        if not row:
            return None
        return {
            "step": row["step"],
            "temp_name": row["temp_name"],
            "temp_username": row["temp_username"],
            "temp_target": row["temp_target"],
            "temp_remind_text": row["temp_remind_text"]
        }

def set_user_state(user_id: int, step: str, temp_name: str = None, temp_username: str = None, temp_target: str = None, temp_remind_text: str = None):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO user_states (user_id, step, temp_name, temp_username, temp_target, temp_remind_text) VALUES (?, ?, ?, ?, ?, ?) "
            "ON CONFLICT(user_id) DO UPDATE SET step = excluded.step, temp_name = excluded.temp_name, temp_username = excluded.temp_username, temp_target = excluded.temp_target, temp_remind_text = excluded.temp_remind_text",
            (user_id, step, temp_name, temp_username, temp_target, temp_remind_text)
        )

def clear_user_state(user_id: int):
    with get_db() as conn:
        conn.execute("DELETE FROM user_states WHERE user_id = ?", (user_id,))

def find_chat_id_by_target(target_str: str):
    target_clean = (target_str or "").lstrip("@").strip().lower()
    
    # 1. Проверяем users.txt: вдруг ввели имя, а в файле есть username
    mapped_username = None
    if os.path.exists(USERS_FILE):
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            for line in f:
                parts = [p.strip() for p in line.split(":", 2)]
                if len(parts) >= 2:
                    name = parts[0].strip().lower()
                    uname = parts[1].lstrip("@").strip().lower()
                    if target_clean in (name, uname):
                        mapped_username = uname
                        break

    search_targets = [target_clean]
    if mapped_username:
        search_targets.append(mapped_username)

    # 2. Ищем chat_id в базе зарегистрированных пользователей
    with get_db() as conn:
        for t in search_targets:
            row = conn.execute(
                "SELECT user_id FROM telegram_users WHERE LOWER(username) = ? OR LOWER(first_name) = ? LIMIT 1",
                (t, t)
            ).fetchone()
            if row:
                return row["user_id"]
    return None


def get_user_mode(user_id: int, is_admin: bool) -> str:
    default_mode = "admin" if is_admin else "worker"
    with get_db() as conn:
        row = conn.execute("SELECT mode FROM user_settings WHERE user_id = ?", (user_id,)).fetchone()
        return row["mode"] if row else default_mode


def set_user_mode(user_id: int, mode: str):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO user_settings (user_id, mode) VALUES (?, ?) "
            "ON CONFLICT(user_id) DO UPDATE SET mode = excluded.mode",
            (user_id, mode)
        )


def save_chat_message(user_id: int, role: str, content: str):
    with get_db() as conn:
        conn.execute(
            "INSERT INTO chat_history (user_id, role, content, created_at) VALUES (?, ?, ?, datetime('now', 'localtime'))",
            (user_id, role, content)
        )


def get_recent_history(user_id: int, limit: int = 10):
    with get_db() as conn:
        rows = conn.execute(
            "SELECT role, content FROM chat_history WHERE user_id = ? ORDER BY id DESC LIMIT ?",
            (user_id, limit)
        ).fetchall()
        return [{"role": r["role"], "content": r["content"]} for r in reversed(rows)]


def clear_user_history(user_id: int):
    with get_db() as conn:
        conn.execute("DELETE FROM chat_history WHERE user_id = ?", (user_id,))

def add_user_rule(user_id: int, rule_text: str) -> bool:
    with get_db() as conn:
        cur = conn.execute("INSERT INTO user_rules (user_id, rule_text) VALUES (?, ?)", (user_id, rule_text.strip()))
        return cur.rowcount > 0

def get_user_rules(user_id: int) -> list:
    with get_db() as conn:
        rows = conn.execute("SELECT id, rule_text FROM user_rules WHERE user_id = ? ORDER BY id ASC", (user_id,)).fetchall()
        return [{"id": r["id"], "text": r["rule_text"]} for r in rows]

def clear_user_rules(user_id: int) -> int:
    with get_db() as conn:
        cur = conn.execute("DELETE FROM user_rules WHERE user_id = ?", (user_id,))
        return cur.rowcount


def cancel_waiting_edits(user_id: int):
    with get_db() as conn:
        conn.execute("UPDATE drafts SET waiting_edit = 0 WHERE user_id = ?", (user_id,))


def create_draft(user_id: int, uinfo: dict, parsed_data: dict) -> int:
    with get_db() as conn:
        conn.execute("UPDATE drafts SET waiting_edit = 0 WHERE user_id = ?", (user_id,))
        cur = conn.execute(
            "INSERT INTO drafts (user_id, uinfo_json, data_json, waiting_edit, created_at) "
            "VALUES (?, ?, ?, 0, datetime('now', 'localtime'))",
            (user_id, json.dumps(uinfo, ensure_ascii=False), json.dumps(parsed_data, ensure_ascii=False))
        )
        return cur.lastrowid


def get_draft(draft_id: int):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM drafts WHERE id = ?", (draft_id,)).fetchone()
        if not row:
            return None
        return {
            "id": row["id"],
            "user_id": row["user_id"],
            "uinfo": json.loads(row["uinfo_json"]),
            "data": json.loads(row["data_json"]),
            "waiting_edit": bool(row["waiting_edit"])
        }


def get_waiting_edit_draft(user_id: int):
    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM drafts WHERE user_id = ? AND waiting_edit = 1 ORDER BY id DESC LIMIT 1",
            (user_id,)
        ).fetchone()
        if not row:
            return None
        return {
            "id": row["id"],
            "user_id": row["user_id"],
            "uinfo": json.loads(row["uinfo_json"]),
            "data": json.loads(row["data_json"]),
            "waiting_edit": True
        }


def update_draft(draft_id: int, parsed_data: dict = None, waiting_edit: bool = False):
    with get_db() as conn:
        if parsed_data is not None:
            conn.execute(
                "UPDATE drafts SET data_json = ?, waiting_edit = ? WHERE id = ?",
                (json.dumps(parsed_data, ensure_ascii=False), int(waiting_edit), draft_id)
            )
        else:
            conn.execute("UPDATE drafts SET waiting_edit = ? WHERE id = ?", (int(waiting_edit), draft_id))


def delete_draft(draft_id: int):
    with get_db() as conn:
        conn.execute("DELETE FROM drafts WHERE id = ?", (draft_id,))


def get_admin_keyboard():
    kb = [
        [KeyboardButton("📊 Финансовая сводка"), KeyboardButton("📥 Выгрузить Excel")],
        [KeyboardButton("📋 Последние записи"), KeyboardButton("📩 Создать Сообщение")],
        [KeyboardButton("⚙️ Управление персоналом"), KeyboardButton("🔄 Сбросить контекст")],
        [KeyboardButton("👷 Режим Сотрудника")]
    ]
    return ReplyKeyboardMarkup(kb, resize_keyboard=True, is_persistent=True)


def get_worker_keyboard(is_admin: bool = False):
    kb = [[KeyboardButton("📩 Создать Сообщение")]]
    if is_admin:
        kb.append([KeyboardButton("👑 Режим Администратора")])
    return ReplyKeyboardMarkup(kb, resize_keyboard=True, is_persistent=True)


def get_draft_keyboard(draft_id: int):
    return InlineKeyboardMarkup([
        [
            InlineKeyboardButton("✅ Сохранить в базу", callback_data=f"confirm_ok:{draft_id}"),
            InlineKeyboardButton("✏️ Исправить", callback_data=f"confirm_edit:{draft_id}")
        ],
        [InlineKeyboardButton("❌ Отмена", callback_data=f"confirm_cancel:{draft_id}")]
    ])


async def safe_reply(message, text: str, reply_markup=None):
    try:
        return await message.reply_text(text, parse_mode="Markdown", reply_markup=reply_markup)
    except BadRequest:
        return await message.reply_text(text, reply_markup=reply_markup)


async def safe_edit(query, text: str, reply_markup=None):
    try:
        return await query.edit_message_text(text, parse_mode="Markdown", reply_markup=reply_markup)
    except BadRequest:
        return await query.edit_message_text(text, reply_markup=reply_markup)


async def send_long_message(message, text: str, max_len: int = 4000):
    if len(text) <= max_len:
        await safe_reply(message, text)
        return

    chunks = []
    current_chunk = ""
    for line in text.split("\n"):
        if len(current_chunk) + len(line) + 1 > max_len:
            if current_chunk:
                chunks.append(current_chunk)
                current_chunk = line
            else:
                chunks.append(line[:max_len])
                current_chunk = line[max_len:]
        else:
            current_chunk = f"{current_chunk}\n{line}" if current_chunk else line

    if current_chunk:
        chunks.append(current_chunk)

    for chunk in chunks:
        await safe_reply(message, chunk)


def decode_text_bytes(file_bytes: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return file_bytes.decode(enc)[:20000]
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")[:20000]


def parse_excel_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    rows = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows.append(f"=== Лист: {sheet} ===")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                rows.append(" | ".join(str(c) if c is not None else "" for c in row).strip())
    return "\n".join(rows)[:20000]


def parse_excel_xls(file_bytes: bytes) -> str:
    wb = xlrd.open_workbook(file_contents=file_bytes)
    rows = []
    for sheet_name in wb.sheet_names():
        sheet = wb.sheet_by_name(sheet_name)
        rows.append(f"=== Лист: {sheet_name} ===")
        for row_idx in range(sheet.nrows):
            row_vals = sheet.row_values(row_idx)
            if any(v != "" and v is not None for v in row_vals):
                rows.append(" | ".join(str(v) if v != "" else "" for v in row_vals).strip())
    return "\n".join(rows)[:20000]


def parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text = []
    for idx, page in enumerate(reader.pages):
        t = page.extract_text()
        if t:
            text.append(f"--- Стр {idx+1} ---\n{t}")
    return "\n".join(text)[:20000]


def generate_excel_bytes() -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Операции"
    ws1.append(["ID", "Дата записи", "Сотрудник", "Дата док.", "Тип", "Юрлицо", "Контрагент", "Вид", "Сумма (₸)", "Детали"])

    ws2 = wb.create_sheet(title="Сводки")
    ws2.append(["ID", "Дата записи", "Сотрудник", "Заголовок", "Описание карточки"])

    with get_db() as conn:
        rows = conn.execute(
            "SELECT id, created_at, added_by_name, doc_date, doc_type, company, counterparty, operation_type, amount, details "
            "FROM transactions ORDER BY id DESC"
        ).fetchall()
        for r in rows:
            ws1.append(list(r))

        sum_rows = conn.execute(
            "SELECT id, created_at, added_by_name, title, formatted_summary FROM summaries ORDER BY id DESC"
        ).fetchall()
        for sr in sum_rows:
            ws2.append(list(sr))

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


# Нативный инструмент веб-поиска Claude API
WEB_SEARCH_TOOL = {
    "type": "web_search_20250305",
    "name": "web_search"
}


PARSE_DATE_TOOL = {
    "name": "parse_datetime",
    "description": "Парсит текстовое описание даты и времени в точный формат YYYY-MM-DD HH:MM или определяет флаг немедленной отправки.",
    "input_schema": {
        "type": "object",
        "properties": {
            "send_now": {
                "type": "boolean",
                "description": "True, если пользователь хочет отправить Сообщение сразу / прямо сейчас."
            },
            "scheduled_datetime": {
                "type": "string",
                "description": "Дата и время в формате 'YYYY-MM-DD HH:MM'. Если send_now=True, может быть текущей датой."
            },
            "readable_time": {
                "type": "string",
                "description": "Понятное человеку описание срока (например, 'через 15 минут', 'завтра в 14:00', 'прямо сейчас')."
            }
        },
        "required": ["send_now", "scheduled_datetime", "readable_time"]
    }
}


async def parse_reminder_time_via_claude(user_input: str) -> dict:
    now_dt = get_current_almaty_time()
    now_str = now_dt.strftime("%Y-%m-%d %H:%M (%A, Asia/Almaty UTC+5)")
    system_prompt = (
        f"Ты — парсер даты и времени. Текущее время сервера: {now_str}.\n"
        "Переведи фразу пользователя о сроке (например 'через 2 часа', 'завтра в 10 утра', 'прямо сейчас', 'сегодня в 18:00') "
        "в точную целевую дату в формате YYYY-MM-DD HH:MM.\n"
        "Обязательно вызови инструмент parse_datetime."
    )
    resp = await claude_client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=500,
        system=system_prompt,
        tools=[PARSE_DATE_TOOL],
        tool_choice={"type": "tool", "name": "parse_datetime"},
        messages=[{"role": "user", "content": user_input}]
    )
    for block in resp.content:
        if block.type == "tool_use" and block.name == "parse_datetime":
            return block.input
    return {"send_now": False, "scheduled_datetime": now_dt.strftime("%Y-%m-%d %H:%M"), "readable_time": user_input}


async def reminder_scheduler(app: Application):
    while True:
        try:
            now_str = get_current_almaty_time().strftime("%Y-%m-%d %H:%M")
            with get_db() as conn:
                due_records = conn.execute(
                    "SELECT id, creator_name, target_user, remind_text, remind_time FROM reminders WHERE is_sent = 0 AND remind_time <= ?",
                    (now_str,)
                ).fetchall()

            for rec in due_records:
                target_chat_id = find_chat_id_by_target(rec["target_user"])

                alert_text = (
                    f"📩 **Сообщение!**\n\n"
                    f"👤 **Кому:** {rec['target_user']}\n"
                    f"📝 **Суть:** {rec['remind_text']}\n"
                    f"✍️ **От кого:** {rec['creator_name']}"
                )

                # Отправляем адресату, а если он не найден — создателю с пометкой
                if target_chat_id:
                    try:
                        await app.bot.send_message(chat_id=target_chat_id, text=alert_text, parse_mode="Markdown")
                    except Exception as err:
                        logger.error(f"Failed to send reminder to target #{rec['id']}: {err}")
                else:
                    try:
                        await app.bot.send_message(
                            chat_id=rec["creator_id"],
                            text=f"⚠️ Не удалось доставить адресату **{rec['target_user']}** (пользователь не запустил бота):\n\n{alert_text}",
                            parse_mode="Markdown"
                        )
                    except Exception as err:
                        logger.error(f"Failed to notify creator #{rec['id']}: {err}")

                with get_db() as conn:
                    conn.execute("UPDATE reminders SET is_sent = 1 WHERE id = ?", (rec["id"],))

        except Exception as e:
            logger.error(f"Reminder scheduler loop error: {e}")

        await asyncio.sleep(30)


EXTRACTION_TOOL = {
    "name": "record_financial_document",
    "description": "Извлекает структурированные финансовые операции из документа, чека, таблицы или текста.",
    "input_schema": {
        "type": "object",
        "properties": {
            "readable_card": {
                "type": "string",
                "description": "Краткая, наглядная сводка операции с суммами, датами, контрагентом и эмодзи для показа в Telegram."
            },
            "summary_title": {
                "type": "string",
                "description": "Короткое название операции или документа."
            },
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "doc_date": {"type": "string", "description": "Дата документа в формате ДД.ММ.ГГГГ"},
                        "doc_type": {"type": "string", "description": "Чек / Счет / Выписка / Акт / Корректировка / Остаток"},
                        "company": {"type": "string", "description": "AS GROUP / ИП БУХLINE / АСТ / АСБ / Kaspi / Свободные средства"},
                        "counterparty": {"type": "string", "description": "Контрагент или основание операции"},
                        "operation_type": {
                            "type": "string",
                            "enum": ["Приход", "Расход", "Свободные средства", "Остаток", "Корректировка"]
                        },
                        "amount": {"type": "number", "description": "Сумма операции числом (только положительное число)"},
                        "details": {"type": "string", "description": "Назначение платежа или детали"}
                    },
                    "required": ["doc_date", "doc_type", "company", "counterparty", "operation_type", "amount", "details"]
                }
            }
        },
        "required": ["readable_card", "summary_title", "items"]
    }
}

SYSTEM_EXTRACTOR = """Ты — финансовый контролер компании AST GROUP (компании группы: AS GROUP, ИП БУХLINE, АСТ, АСБ).
Проанализируй присланный документ, изображение чека, таблицу или текст.
Обязательно вызови инструмент record_financial_document и заполни все поля максимально точно."""


async def extract_document_via_claude(messages_content: list, previous_draft: dict = None, user_correction: str = None) -> dict:
    today_str = datetime.now().strftime("%d.%m.%Y")
    system_with_date = f"{SYSTEM_EXTRACTOR}\nСегодняшняя дата: {today_str}."

    if previous_draft and user_correction:
        prompt = (
            f"Предыдущие распознанные данные (JSON):\n{json.dumps(previous_draft, ensure_ascii=False)}\n\n"
            f"Правка от пользователя: {user_correction}\n"
            "Пересчитай данные, обнови readable_card и вызови инструмент record_financial_document."
        )
        content = [{"type": "text", "text": prompt}]
    else:
        content = messages_content

    resp = await claude_client.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=3000,
        system=system_with_date,
        tools=[EXTRACTION_TOOL],
        tool_choice={"type": "tool", "name": "record_financial_document"},
        messages=[{"role": "user", "content": content}]
    )

    for block in resp.content:
        if block.type == "tool_use" and block.name == "record_financial_document":
            return block.input

    raise ValueError("Модель не вернула ожидаемую структуру данных.")


async def ask_financial_analyst(user_id: int, question: str) -> str:
    # Проверяем, требует ли запрос данных финансовой базы
    financial_keywords = ["баланс", "сводка", "деньги", "транзакц", "остаток", "расход", "приход", "касса", "отчет", "выручка", "kzt", "тенге"]
    needs_finance_data = any(kw in question.lower() for kw in financial_keywords)

    context_str = ""
    if needs_finance_data:
        with get_db() as conn:
            total_count = conn.execute("SELECT COUNT(*) as cnt FROM transactions").fetchone()["cnt"]
            totals_by_type = conn.execute(
                "SELECT operation_type, SUM(amount) as total FROM transactions GROUP BY operation_type"
            ).fetchall()
            txs = conn.execute(
                "SELECT id, doc_date, doc_type, company, counterparty, operation_type, amount, details "
                "FROM transactions ORDER BY id DESC LIMIT 100"
            ).fetchall()

        sum_initial = sum(float(r["total"] or 0.0) for r in totals_by_type if r["operation_type"] in ("Остаток", "Свободные средства"))
        sum_income = sum(float(r["total"] or 0.0) for r in totals_by_type if r["operation_type"] == "Приход")
        sum_expense = sum(float(r["total"] or 0.0) for r in totals_by_type if r["operation_type"] == "Расход")
        exact_balance = (sum_initial + sum_income) - sum_expense

        context_str = (
            f"ФИНАНСОВЫЙ КОНТЕКСТ БАЗЫ (Всего записей: {total_count}):\n"
            f"Баланс: {exact_balance:,.0f} KZT | Приход: {sum_income:,.0f} KZT | Расход: {sum_expense:,.0f} KZT\n"
            "ПОСЛЕДНИЕ ОПЕРАЦИИ:\n"
        )
        for t in reversed(txs):
            context_str += f"- #{t['id']} | {t['doc_date']} | {t['company']} -> {t['counterparty']} | {t['operation_type']}: {t['amount']:,.0f} KZT ({t['details']})\n"
        context_str += "\n"

    # Подтягиваем сохранённые правила пользователя
    user_rules = get_user_rules(user_id)
    rules_block = ""
    if user_rules:
        rules_block = "ПОЛЬЗОВАТЕЛЬСКИЕ ПРАВИЛА:\n" + "\n".join(f"- {r['text']}" for r in user_rules) + "\n\n"

    now_almaty = get_current_almaty_time()
    today_str = now_almaty.strftime("%d.%m.%Y %H:%M (Asia/Almaty, UTC+5)")
    system_prompt = (
        f"Ты — корпоративный ассистент компании AST GROUP. Системное время: {today_str}.\n"
        f"{rules_block}"
        "Правила ответа:\n"
        "1. Не смешивай контекст разных задач. Если запрос не касается финансов, не упоминай счета и балансы.\n"
        "2. Для внешних данных используй web_search. Если данные не найдены, прямо скажи об этом, не выдумывай.\n"
        "3. При расчётах проверяй формулы и явно показывай контрольную сумму."
    )

    history = get_recent_history(user_id, limit=6)
    messages_payload = [{"role": h["role"], "content": h["content"]} for h in history]
    messages_payload.append({"role": "user", "content": f"{context_str}\n\nЗапрос: {question}"})

    current_messages = list(messages_payload)

    for _ in range(3):
        resp = await claude_client.messages.create(
            model=CLAUDE_MODEL,
            max_tokens=3000,
            system=system_prompt,
            tools=[WEB_SEARCH_TOOL],
            messages=current_messages
        )

        if resp.stop_reason == "tool_use":
            tool_results = []
            for block in resp.content:
                if block.type == "tool_use":
                    # Локальный вызов только для погоды
                    if block.name == "get_weather_forecast":
                        target_city = block.input.get("city", "Астана")
                        res_text = await asyncio.to_thread(get_weather_forecast, target_city)
                        tool_results.append({
                            "type": "tool_result",
                            "tool_use_id": block.id,
                            "content": str(res_text)
                        })

            if not tool_results:
                break

            current_messages.append({"role": "assistant", "content": resp.content})
            current_messages.append({"role": "user", "content": tool_results})
        else:
            break

    texts = [b.text for b in resp.content if getattr(b, "type", "") == "text" and hasattr(b, "text")]
    answer = "\n".join(texts).strip()

    if not answer:
        answer = "Не удалось сформировать ответ. Попробуйте уточнить запрос."

    save_chat_message(user_id, "user", question)
    save_chat_message(user_id, "assistant", answer)
    return answer


async def start_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    uinfo = get_user_info(user.username)
    if not uinfo:
        await safe_reply(update.message, f"⛔ Доступ запрещен. ID: {user.id}")
        return

    cancel_waiting_edits(user.id)
    
    # Сохраняем пользователя для доставки напоминаний
    with get_db() as conn:
        conn.execute(
            "INSERT INTO telegram_users (user_id, username, first_name) VALUES (?, ?, ?) "
            "ON CONFLICT(user_id) DO UPDATE SET username = excluded.username, first_name = excluded.first_name",
            (user.id, (user.username or "").lower(), uinfo["name"])
        )

    mode = "admin" if uinfo["is_admin"] else "worker"
    set_user_mode(user.id, mode)
    kb = get_admin_keyboard() if uinfo["is_admin"] else get_worker_keyboard(is_admin=False)

    await safe_reply(
        update.message,
        f"👋 Здравствуйте, {uinfo['name']}!\n"
        f"Режим: **{'Администратор' if uinfo['is_admin'] else 'Сотрудник'}**.",
        reply_markup=kb
    )


async def delete_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    uinfo = get_user_info(user.username)
    if not uinfo or not uinfo["is_admin"]:
        return

    if not context.args:
        await safe_reply(update.message, "ℹ️ Укажите ID записи. Пример: /del 3")
        return

    try:
        tx_id = int(context.args[0])
    except ValueError:
        return

    with get_db() as conn:
        record = conn.execute("SELECT id, company, counterparty, amount FROM transactions WHERE id = ?", (tx_id,)).fetchone()
        if not record:
            await safe_reply(update.message, f"⚠️ Запись #{tx_id} не найдена.")
            return
        conn.execute("DELETE FROM transactions WHERE id = ?", (tx_id,))

    await safe_reply(update.message, f"🗑 Запись #{tx_id} удалена.")


async def send_last_records(update: Update):
    with get_db() as conn:
        rows = conn.execute("SELECT id, created_at, company, counterparty, amount, details FROM transactions ORDER BY id DESC LIMIT 5").fetchall()

    if not rows:
        await update.message.reply_text("В базе пока нет записей.")
        return

    await safe_reply(update.message, "📋 **Последние записи:**")
    for r in rows:
        msg = f"🔹 **ID #{r['id']}**\n• {r['company']} -> {r['counterparty']}\n• Сумма: **{r['amount']:,.0f} ₸**"
        del_kb = InlineKeyboardMarkup([[InlineKeyboardButton(f"🗑 Удалить #{r['id']}", callback_data=f"del_tx_{r['id']}")]])
        await safe_reply(update.message, msg, reply_markup=del_kb)


async def callback_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user = query.from_user
    uinfo = get_user_info(user.username)
    if not uinfo or not uinfo["is_admin"]:
        return

    data = query.data

    if data == "staff_menu":
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("👥 Список сотрудников", callback_data="staff_list")],
            [InlineKeyboardButton("➕ Добавить сотрудника", callback_data="staff_add")],
            [InlineKeyboardButton("❌ Удалить сотрудника", callback_data="staff_delete_select")],
            [InlineKeyboardButton("🔙 Закрыть", callback_data="cancel_action")]
        ])
        await safe_edit(query, "⚙️ **Управление персоналом**\nВыберите необходимое действие:", reply_markup=kb)
        return

    if data == "staff_list":
        content = "Файл пуст."
        if os.path.exists(USERS_FILE):
            with open(USERS_FILE, "r", encoding="utf-8") as f:
                content = f.read()
        kb = InlineKeyboardMarkup([[InlineKeyboardButton("🔙 Назад", callback_data="staff_menu")]])
        await safe_edit(query, f"👥 **Список сотрудников:**\n```text\n{content}\n```", reply_markup=kb)
        return

    if data == "staff_add":
        set_user_state(user.id, step="awaiting_name")
        await safe_edit(query, "👤 Введите **Имя** нового сотрудника (например: Айдар):")
        return

    if data == "staff_delete_select":
        inline_kb = []
        if os.path.exists(USERS_FILE):
            with open(USERS_FILE, "r", encoding="utf-8") as f:
                for line in f:
                    stripped = line.strip()
                    if not stripped or stripped.startswith("#") or ":" not in stripped:
                        continue
                    parts = [p.strip() for p in stripped.split(":", 2)]
                    name, uname = parts[0], parts[1]
                    if uname.lstrip("@").lower() in ["lexaengineer1441", "ast_aidar"]:
                        continue
                    inline_kb.append([InlineKeyboardButton(f"🗑 {name} ({uname})", callback_data=f"del_user_{uname}")])
        inline_kb.append([InlineKeyboardButton("🔙 Назад", callback_data="staff_menu")])
        await safe_edit(query, "👥 **Выберите сотрудника для удаления:**", reply_markup=InlineKeyboardMarkup(inline_kb))
        return

    if data.startswith("del_user_"):
        target_uname = data.replace("del_user_", "").lstrip("@").strip().lower()
        lines = []
        found = False
        deleted_line = ""
        
        if os.path.exists(USERS_FILE):
            with open(USERS_FILE, "r", encoding="utf-8") as f:
                for line in f:
                    stripped = line.strip()
                    if not stripped or stripped.startswith("#") or ":" not in stripped:
                        lines.append(line)
                        continue
                    parts = [p.strip() for p in stripped.split(":", 2)]
                    file_uname = parts[1].lstrip("@").strip().lower()
                    if file_uname == target_uname:
                        found = True
                        deleted_line = stripped
                    else:
                        lines.append(line)

        if found:
            with open(USERS_FILE, "w", encoding="utf-8") as f:
                f.writelines(lines)
            await safe_edit(query, f"✅ **Сотрудник удален!**\n{deleted_line}")
        else:
            await safe_edit(query, f"⚠️ Пользователь @{target_uname} не найден.")
        return

    if data.startswith("del_tx_"):
        tx_id = int(data.replace("del_tx_", ""))
        with get_db() as conn:
            conn.execute("DELETE FROM transactions WHERE id = ?", (tx_id,))
        await safe_edit(query, f"🗑 Запись **ID #{tx_id}** удалена.")
        return

    if data == "cancel_action":
        await safe_edit(query, "❌ Отменено.")
        return

    if ":" in data:
        action, draft_id_str = data.split(":", 1)
        draft_id = int(draft_id_str)
    else:
        return

    draft = get_draft(draft_id)
    if not draft:
        await safe_edit(query, "Черновик устарел.")
        return

    if action == "confirm_ok":
        parsed = draft["data"]
        draft_uinfo = draft["uinfo"]
        items = parsed.get("items", [])

        with get_db() as conn:
            conn.execute(
                "INSERT INTO summaries (added_by_name, title, formatted_summary, raw_json, created_at) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))",
                (draft_uinfo["name"], parsed.get("summary_title", "Отчет"), parsed.get("readable_card", ""), json.dumps(parsed, ensure_ascii=False))
            )
            for item in items:
                conn.execute("""
                    INSERT INTO transactions 
                    (added_by_username, added_by_name, doc_date, doc_type, company, counterparty, operation_type, amount, details, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now', 'localtime'))
                """, (
                    draft_uinfo["username"], draft_uinfo["name"],
                    item.get("doc_date", ""), item.get("doc_type", ""),
                    item.get("company", ""), item.get("counterparty", ""),
                    item.get("operation_type", ""), float(item.get("amount") or 0),
                    item.get("details", "")
                ))
            conn.execute("DELETE FROM drafts WHERE id = ?", (draft_id,))

        await safe_edit(query, f"✅ **Успешно записано в базу данных SQLite!**\n\n• Таблица: `transactions` ({len(items)} записей)\n• Таблица: `summaries` (1 сводка)\n\n{parsed.get('readable_card', '')}")

    elif action == "confirm_edit":
        cancel_waiting_edits(user.id)
        update_draft(draft_id, waiting_edit=True)
        await safe_reply(query.message, "✏️ Напишите исправление текстом:")

    elif action == "confirm_cancel":
        delete_draft(draft_id)
        await safe_edit(query, "❌ Отменено.")


async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    uinfo = get_user_info(user.username)
    if not uinfo:
        return

    text = update.message.text.strip()

    # Проверка активных состояний (добавление сотрудника или создание напоминания)
    state = get_user_state(user.id)
    back_kb = get_admin_keyboard() if uinfo["is_admin"] else get_worker_keyboard(is_admin=False)

    if state:
        if text.lower() in ["отмена", "/cancel"]:
            clear_user_state(user.id)
            await safe_reply(update.message, "❌ Действие отменено.", reply_markup=back_kb)
            return

        # Шаги создания сотрудника
        if uinfo["is_admin"] and state["step"] == "awaiting_name":
            set_user_state(user.id, step="awaiting_username", temp_name=text)
            await safe_reply(update.message, "🔗 Введите **Telegram Username** сотрудника (например: @username):")
            return

        elif uinfo["is_admin"] and state["step"] == "awaiting_username":
            uname = text if text.startswith("@") else f"@{text}"
            set_user_state(user.id, step="awaiting_role", temp_name=state["temp_name"], temp_username=uname)
            await safe_reply(update.message, "📝 Введите **Роль / Описание** сотрудника (например: Менеджер):")
            return

        elif uinfo["is_admin"] and state["step"] == "awaiting_role":
            name, uname, role = state["temp_name"], state["temp_username"], text
            clear_user_state(user.id)
            new_line = f"{name} : {uname} : {role}"
            with open(USERS_FILE, "a", encoding="utf-8") as f:
                f.write(f"\n{new_line}")
            await safe_reply(update.message, f"✅ **Сотрудник успешно добавлен!**\n{new_line}", reply_markup=back_kb)
            return

        # Шаги создания напоминания
        elif state["step"] == "remind_target":
            cleaned_target = text.strip().lower()
            if cleaned_target in ["мне", "себе", "я", "сам", "сама"]:
                target_value = user.username or uinfo["name"]
            else:
                target_value = text

            set_user_state(user.id, step="remind_text", temp_target=target_value)
            await safe_reply(update.message, f"📝 Что именно нужно напомнить для **{target_value}**?")
            return

        elif state["step"] == "remind_text":
            set_user_state(user.id, step="remind_time", temp_target=state["temp_target"], temp_remind_text=text)
            await safe_reply(
                update.message,
                "🗓 Когда отправить Сообщение?\n"
                "Напишите в свободной форме (например: `прямо сейчас`, `завтра в 11:30`, `через 45 минут`, `в пятницу в 18:00`):"
            )
            return

        elif state["step"] == "remind_time":
            target = state["temp_target"]
            remind_msg = state["temp_remind_text"]
            clear_user_state(user.id)

            await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
            parsed_dt = await parse_reminder_time_via_claude(text)
            target_chat_id = find_chat_id_by_target(target)

            alert_text = (
                f"📩 **Сообщение!**\n\n"
                f"👤 **Кому:** {target}\n"
                f"📝 **Суть:** {remind_msg}\n"
                f"✍️ **От кого:** {uinfo['name']}"
            )

            if parsed_dt.get("send_now"):
                sent_to_target = False
                if target_chat_id:
                    try:
                        await context.bot.send_message(chat_id=target_chat_id, text=alert_text, parse_mode="Markdown")
                        sent_to_target = True
                    except Exception as e:
                        logger.error(f"Не удалось отправить пользователю: {e}")

                if sent_to_target:
                    await safe_reply(update.message, f"⚡️ Срочное Сообщение отправлено пользователю **{target}**!", reply_markup=back_kb)
                else:
                    await safe_reply(
                        update.message,
                        f"⚠️ Пользователь **{target}** еще не запускал бота (или не найден в базе).\nСообщение отображено здесь:\n\n{alert_text}",
                        reply_markup=back_kb
                    )
                return

            sched_time = parsed_dt.get("scheduled_datetime")
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO reminders (creator_id, creator_name, target_user, remind_text, remind_time) VALUES (?, ?, ?, ?, ?)",
                    (user.id, uinfo["name"], target, remind_msg, sched_time)
                )

            status_note = "" if target_chat_id else "\n⚠️ _(Внимание: сотрудник еще ни разу не нажал /start в боте, Сообщение придет вам, если он не активирует бота)_"
            await safe_reply(
                update.message,
                f"✅ **Сообщение запланировано!**\n\n"
                f"👤 **Кому:** {target}\n"
                f"📝 **Суть:** {remind_msg}\n"
                f"📩 **Срок:** {sched_time} ({parsed_dt.get('readable_time', '')}){status_note}",
                reply_markup=back_kb
            )
            return

    # Обработка команд изоляции контекста и правил
    clean_cmd = text.strip().lower()
    if clean_cmd in ["новая задача", "/new", "сбросить контекст"]:
        cancel_waiting_edits(user.id)
        clear_user_history(user.id)
        clear_user_state(user.id)
        await safe_reply(update.message, "🆕 **Контекст очищен.** Начата новая задача без привязки к предыдущим диалогам.", reply_markup=back_kb)
        return

    if clean_cmd in ["отменить", "отмена"]:
        cancel_waiting_edits(user.id)
        clear_user_state(user.id)
        await safe_reply(update.message, "❌ Текущее действие отменено.", reply_markup=back_kb)
        return

    if clean_cmd.startswith("сохранить правило:") or clean_cmd.startswith("запомни:"):
        rule_content = text.split(":", 1)[1].strip()
        if rule_content:
            success = add_user_rule(user.id, rule_content)
            if success:
                await safe_reply(update.message, f"💾 **Успешно сохранено в БД (таблица user_rules):**\n«{rule_content}»\nПравило будет учитываться во всех будущих ответах.")
            else:
                await safe_reply(update.message, "⚠️ Ошибка: запись в базу данных не удалась.")
        else:
            await safe_reply(update.message, "ℹ️ Укажите текст правила после двоеточия. Пример:\n`Сохранить правило: Все суммы округлять до целых`")
        return

    if clean_cmd in ["мои правила", "правила"]:
        rules = get_user_rules(user.id)
        if not rules:
            await safe_reply(update.message, "У вас нет сохранённых персональных правил.\nДобавить: `Сохранить правило: <текст>`")
        else:
            formatted_rules = "\n".join([f"{idx+1}. {r['text']} (ID: {r['id']})" for idx, r in enumerate(rules)])
            await safe_reply(update.message, f"📋 **Ваши сохранённые правила:**\n{formatted_rules}\n\nОчистить: `Очистить правила`")
        return

    if clean_cmd in ["очистить правила"]:
        cnt = clear_user_rules(user.id)
        await safe_reply(update.message, f"🗑 Удалено сохранённых правил: {cnt}.")
        return

    if text == "📩 Создать Сообщение":
        cancel_waiting_edits(user.id)
        set_user_state(user.id, step="remind_target")
        await safe_reply(update.message, "👤 Кому адресовано Сообщение? Введите имя или @username:")
        return


    if uinfo["is_admin"]:
        if text in ["👷 Режим Сотрудника"]:
            cancel_waiting_edits(user.id)
            set_user_mode(user.id, "worker")
            await update.message.reply_text("👷 Включен режим сотрудника.", reply_markup=get_worker_keyboard(is_admin=True))
            return

        elif text in ["👑 Режим Администратора"]:
            cancel_waiting_edits(user.id)
            set_user_mode(user.id, "admin")
            await update.message.reply_text("👑 Включен режим администратора.", reply_markup=get_admin_keyboard())
            return

        elif text == "🔄 Сбросить контекст":
            cancel_waiting_edits(user.id)
            clear_user_history(user.id)
            await update.message.reply_text("🔄 История очищена.", reply_markup=get_admin_keyboard())
            return

        elif text == "📋 Последние записи":
            cancel_waiting_edits(user.id)
            await send_last_records(update)
            return

        elif text == "📊 Финансовая сводка":
            cancel_waiting_edits(user.id)
            await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
            ans = await ask_financial_analyst(user.id, "Сделай полную финансовую сводку.")
            await send_long_message(update.message, ans)
            return

        elif text == "📥 Выгрузить Excel":
            cancel_waiting_edits(user.id)
            bio = await asyncio.to_thread(generate_excel_bytes)
            await update.message.reply_document(document=bio, filename="Финансы_AST_GROUP.xlsx")
            return

        elif text == "⚙️ Управление персоналом":
            cancel_waiting_edits(user.id)
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton("👥 Список сотрудников", callback_data="staff_list")],
                [InlineKeyboardButton("➕ Добавить сотрудника", callback_data="staff_add")],
                [InlineKeyboardButton("❌ Удалить сотрудника", callback_data="staff_delete_select")],
                [InlineKeyboardButton("🔙 Закрыть", callback_data="cancel_action")]
            ])
            await safe_reply(update.message, "⚙️ **Управление персоналом**\nВыберите действие:", reply_markup=kb)
            return

    waiting_draft = get_waiting_edit_draft(user.id)
    if waiting_draft:
        await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
        try:
            updated = await extract_document_via_claude([], previous_draft=waiting_draft["data"], user_correction=text)
            update_draft(waiting_draft["id"], parsed_data=updated, waiting_edit=False)
            await safe_reply(
                update.message,
                f"📋 **Обновлено:**\n\n{updated.get('readable_card', '')}\n\n❓ **Сохраняем?**",
                reply_markup=get_draft_keyboard(waiting_draft["id"])
            )
        except Exception:
            await update.message.reply_text("⚠️ Не удалось применить правку.")
        return

    mode = get_user_mode(user.id, uinfo["is_admin"])
    if mode == "admin":
        lower_text = text.lower()
        add_triggers = ["добавь", "запиши", "внеси", "зафиксируй", "сохрани", "свободных"]
        if any(w in lower_text for w in add_triggers) and not text.startswith("📊"):
            await process_incoming_doc(update, context, uinfo, [{"type": "text", "text": f"Фиксация операции: {text}"}])
            return

        await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
        ans = await ask_financial_analyst(user.id, text)
        await send_long_message(update.message, ans)
        return

    await process_incoming_doc(update, context, uinfo, [{"type": "text", "text": text}])


async def process_incoming_doc(update: Update, context: ContextTypes.DEFAULT_TYPE, uinfo: dict, content_blocks: list):
    user_id = update.effective_user.id
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
    try:
        parsed = await extract_document_via_claude(content_blocks)
        items = parsed.get("items", [])
        
        # Контрольная проверка арифметики
        total_items_amount = sum(float(item.get("amount") or 0.0) for item in items)
        audit_note = f"\n\n🔍 **Контрольная сверка:** строк в документе: {len(items)}, общая сумма позиций: **{total_items_amount:,.0f} ₸**"
        
        draft_id = create_draft(user_id, uinfo, parsed)
        card_text = f"📋 **Проверьте данные:**\n\n{parsed.get('readable_card', '')}{audit_note}\n\n❓ **Сохранить в базу?**"
        await safe_reply(update.message, card_text, reply_markup=get_draft_keyboard(draft_id))
    except Exception as e:
        logger.error(f"Ошибка распознавания данных: {e}", exc_info=True)
        await update.message.reply_text("⚠️ Ошибка распознавания данных. Проверьте формат документа.")


async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    uinfo = get_user_info(user.username)
    if not uinfo:
        return
    photo_file = await update.message.photo[-1].get_file()
    b = bytes(await photo_file.download_as_bytearray())

    # Сохранение фото в папку data
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    uname = (user.username or str(user.id)).lstrip("@")
    file_path = os.path.join(DATA_DIR, f"{timestamp}_{uname}_{photo_file.file_unique_id}.jpg")
    try:
        with open(file_path, "wb") as f:
            f.write(b)
    except Exception as e:
        logger.error(f"Ошибка сохранения фото на диск: {e}")

    b64 = base64.b64encode(b).decode("utf-8")
    caption = update.message.caption or "Распознай чек."
    await process_incoming_doc(update, context, uinfo, [
        {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}},
        {"type": "text", "text": caption}
    ])


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user = update.effective_user
    uinfo = get_user_info(user.username)
    if not uinfo:
        return

    doc = update.message.document
    original_name = doc.file_name or "document"
    fname = original_name.lower()
    fobj = await doc.get_file()
    b = bytes(await fobj.download_as_bytearray())

    # Сохранение документа в папку data
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    uname = (user.username or str(user.id)).lstrip("@")
    clean_original_name = re.sub(r'[^\w\.-]', '_', original_name)
    save_path = os.path.join(DATA_DIR, f"{timestamp}_{uname}_{clean_original_name}")
    try:
        with open(save_path, "wb") as f:
            f.write(b)
    except Exception as e:
        logger.error(f"Ошибка сохранения документа на диск: {e}")

    caption = update.message.caption or f"Распознай документ {doc.file_name}"

    if fname.endswith((".jpg", ".jpeg", ".png", ".webp")):
        b64 = base64.b64encode(b).decode("utf-8")
        await process_incoming_doc(update, context, uinfo, [
            {"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": b64}},
            {"type": "text", "text": caption}
        ])
        return

    try:
        if fname.endswith(".xlsx"):
            extracted = await asyncio.to_thread(parse_excel_xlsx, b)
        elif fname.endswith(".xls"):
            extracted = await asyncio.to_thread(parse_excel_xls, b)
        elif fname.endswith(".pdf"):
            extracted = await asyncio.to_thread(parse_pdf, b)
        elif fname.endswith((".txt", ".csv")):
            extracted = await asyncio.to_thread(decode_text_bytes, b)
        else:
            await update.message.reply_text("Формат не поддерживается.")
            return
    except Exception:
        await update.message.reply_text("⚠️ Не удалось прочитать файл.")
        return

    await process_incoming_doc(update, context, uinfo, [{"type": "text", "text": f"Файл {doc.file_name}:\n{extracted}"}])


def main():
    init_db()
    logger.info("Запуск бота AST GROUP...")
    app = (
        Application.builder()
        .token(TELEGRAM_BOT_TOKEN)
        .read_timeout(90)
        .write_timeout(90)
        .concurrent_updates(True)
        .build()
    )

    app.add_handler(CommandHandler("start", start_command))
    app.add_handler(CommandHandler("del", delete_command))
    app.add_handler(CommandHandler("records", lambda u, c: send_last_records(u)))
    app.add_handler(CallbackQueryHandler(callback_handler))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))

    async def post_init(application: Application):
        asyncio.create_task(reminder_scheduler(application))

    app.post_init = post_init
    app.run_polling(drop_pending_updates=True)


if __name__ == "__main__":
    main()
